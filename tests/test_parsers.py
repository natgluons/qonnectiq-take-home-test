"""Offline tests: synthetic fixtures plus optional integration against local datasets."""
from pathlib import Path

import fitz
import pytest
from docx import Document

from ingest import classify_pdf, run
from parsers import parse_ddr, parse_dgos, parse_glossary


def test_glossary_docx_table(tmp_path):
    p = tmp_path / "Glossaries.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Abbreviation"
    table.rows[0].cells[1].text = "Meaning"
    for a, b in [("A", "A"), ("NPT", "Non-Productive Time – downtime"),
                 ("WIP", "Unknown – (to be confirmed)")]:
        row = table.add_row()
        row.cells[0].text, row.cells[1].text = a, b
    doc.save(p)
    result = parse_glossary(p)
    assert len(result["terms"]) == 2
    assert result["terms"][0]["term"] == "NPT"
    assert result["terms"][1]["needs_confirmation"] is True


def test_synthetic_dgos_without_specific_filename(tmp_path):
    p = tmp_path / "arbitrary_filename.pdf"
    doc = fitz.open()
    page = doc.new_page(width=972, height=1667)
    for x, y, msg in [
        (20, 46, "Current Date : 12-10-2026"),
        (330, 65, "DAILY GEOLOGICAL OPERATIONS SUMMARY"),
        (20, 125, "WELL NAME : TEST-1"),
        (20, 145, "COUNTRY : MYS | MALAYSIA"),
        (900, 95, "73"),
        (20, 250, "Rig up wireline."),
        (20, 310, "Finished drilling."),
        (20, 395, "Perform WL Run #1: TEST."),
        (502, 492, "D12"),
        (825, 492, "2950.00"),
        (720, 552, "14.1"),
        (744, 552, "SBM"),
    ]:
        page.insert_text((x,y), msg, fontsize=10)
    doc.save(p)
    doc.close()
    assert classify_pdf(p) == "DGOS"
    result = parse_dgos(p)
    assert result["report_date"] == "2026-10-12"
    assert result["well_name"] == "TEST-1"
    assert result["current_depth_mddf"] == 2950
    assert "WL Run #1" in result["next_24h_operation"]


def test_synthetic_ddr(tmp_path):
    p = tmp_path / "report.pdf"
    doc = fitz.open()
    page = doc.new_page()
    for i, msg in enumerate([
        "Daily Operation Report",
        "Well: DEMO-1    Report no.: 5    Report date: 12/10/2026",
        "Rig Name: TEST-RIG   Block: PMXXX   Water Depth: 55.00 m",
        "Daily NPT : 1.50 hr   Cumm NPT : 3.50 hr",
        "Daily Cost : 120.00    Cumm Cost : 340.00",
        "Current status : Testing tools.",
        "24 hr summary : Did a test.",
        "24 hr forecast : Repeat the test.",
        "Incident / Accident : No accidents.",
        "Remarks : none",
    ]):
        page.insert_text((25, 50 + 28*i), msg, fontsize=10)
    doc.save(p)
    doc.close()
    assert classify_pdf(p) == "DDR"
    parsed = parse_ddr(p)
    assert parsed["well_name"] == "DEMO-1"
    assert parsed["report_date"] == "2026-10-12"
    assert parsed["cumulative_npt_hours"] == 3.5
    assert parsed["daily_cost_usd"] == 120.0


def test_one_command_ingestion(tmp_path):
    inputs = tmp_path / "datasets"
    inputs.mkdir()
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.add_row().cells[0].text = "NPT"
    table.rows[1].cells[1].text = "Non Productive Time"
    doc.save(inputs / "Glossaries.docx")
    output = tmp_path / "parsed_data"
    result = run(inputs, output)
    assert not result["errors"]
    assert len(result["files_created"]) == 1
    assert (output / "Glossaries.json").is_file()


def test_sqlite_ingestion_is_an_alternative_to_json(tmp_path):
    inputs = tmp_path / "datasets"
    inputs.mkdir()
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.add_row().cells[0].text = "NPT"
    table.rows[1].cells[1].text = "Non Productive Time"
    doc.save(inputs / "Glossaries.docx")
    output = tmp_path / "parsed_data"
    database = output / "corpus.db"
    result = run(inputs, output, database)
    assert not result["errors"]
    assert result["files_created"] == []
    assert result["database"] == str(database)
    assert database.is_file()
    assert not (output / "Glossaries.json").exists()


@pytest.mark.parametrize("suffix,report_type", [
    ("BARAKUDA-1_DGOS_72_20260829.pdf", "DGOS"),
    ("BARAKUDA-1_DGOS_84_20260910.pdf", "DGOS"),
    ("NAGA-2_BARAKUDA-1_DDR_32_19_07_2026_Drill_17_5in_x_20in_Hole.pdf", "DDR"),
])
def test_optional_original_files(suffix, report_type):
    p = Path("datasets") / suffix
    if not p.exists():
        pytest.skip("Private assessment datasets are not bundled in public repo")
    parser = parse_dgos if report_type == "DGOS" else parse_ddr
    result = parser(p)
    assert result["document_type"] == report_type
    assert result["well_name"] == "BARAKUDA-1"
